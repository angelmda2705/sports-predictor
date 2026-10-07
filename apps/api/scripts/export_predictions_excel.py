"""Exporta a Excel las predicciones de los próximos partidos del Mundial 2026.

Usa EXACTAMENTE el mismo cableado que la app (snapshot real + PredictionService
con ratings Elo sembrados de historial internacional). Toda salida es PROBABILÍSTICA
(no certezas) — el Excel incluye el disclaimer.

Uso (desde apps/api):  python -m scripts.export_predictions_excel [--comp WC_2026] [--all]
"""

from __future__ import annotations

import os
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from src.application.backtest_service import BacktestService  # noqa: E402
from src.application.populate_history import populate  # noqa: E402
from src.application.prediction_service import (  # noqa: E402
    PredictionService,
    PredictionUnavailable,
)
from src.domain.catalog import MatchStatus  # noqa: E402
from src.infrastructure.db.engine import build_inmemory_engine, build_session_factory  # noqa: E402
from src.infrastructure.db.prediction_repository import SqlAlchemyPredictionHistory  # noqa: E402
from src.infrastructure.ingestion.snapshot import default_snapshot_path, load_snapshot  # noqa: E402
from src.infrastructure.memory.catalog_repository import InMemoryCatalogRepository  # noqa: E402

MX_TZ = timezone(timedelta(hours=-6))  # México (sin horario de verano desde 2022)

# Estilos
_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_TITLE_FONT = Font(bold=True, size=15, color="1F3864")
_SUB_FONT = Font(italic=True, size=9, color="595959")
_FAV_FILL = PatternFill("solid", fgColor="C6E0B4")  # verde claro: resultado más probable
_BAND_FILL = {
    "alta": PatternFill("solid", fgColor="C6E0B4"),
    "media": PatternFill("solid", fgColor="FFE699"),
    "baja": PatternFill("solid", fgColor="F8CBAD"),
}
_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center")
_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)

_COLS = [
    ("Fecha (México)", 17),
    ("Fase", 22),
    ("Local", 18),
    ("Visitante", 18),
    ("P(Local)", 10),
    ("P(Empate)", 10),
    ("P(Visit.)", 10),
    ("Pronóstico", 20),
    ("Marcador prob.", 13),
    ("Goles L", 8),
    ("Goles V", 8),
    ("Confianza", 11),
    ("Elo L", 8),
    ("Elo V", 8),
    ("Contexto / notas", 42),
]


def _desktop_dir() -> Path:
    home = Path(os.environ.get("USERPROFILE", str(Path.home())))
    for cand in (home / "OneDrive" / "Desktop", home / "Desktop"):
        if cand.is_dir():
            return cand
    return Path.cwd()


def _write_track_record_sheet(ws, catalog, snap, international, predictor) -> None:
    """Segunda hoja: desempeño real del modelo (mismas cifras que /historial)."""
    backtest = BacktestService(
        catalog,
        seed_ratings=snap.seed_ratings,
        seed_counts=snap.seed_counts,
        international_competitions=international,
    )
    engine = build_inmemory_engine()
    history = SqlAlchemyPredictionHistory(build_session_factory(engine))
    history.create_tables(engine)
    populate(catalog, backtest, predictor, history)

    ws.merge_cells("A1:F1")
    ws.cell(1, 1, "Desempeño del modelo (track record)").font = _TITLE_FONT
    ws.merge_cells("A2:F2")
    ws.cell(
        2, 1,
        "Se guardan TODAS las predicciones; al terminar el partido se marca acierto/fallo. "
        "No se ocultan fallos.",
    ).font = _SUB_FONT

    headers = ["Competición", "Guardadas", "Resueltas", "Aciertos 1X2", "Brier", "LogLoss"]
    widths = [24, 12, 12, 14, 10, 10]
    hr = 4
    for j, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(hr, j, h)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = _CENTER
        c.border = _BORDER
        ws.column_dimensions[get_column_letter(j)].width = w

    rows = [
        ("GLOBAL (todas las ligas)", None),
        ("Mundial 2026", "WC_2026"),
        ("Premier League", "ENG_PL"),
    ]
    r = hr + 1
    for label, code in rows:
        tr = history.track_record(competition_code=code)
        vals = [
            label,
            tr.n_predictions,
            tr.n_resolved,
            tr.accuracy if tr.n_resolved else None,
            round(tr.brier, 3) if tr.n_resolved else None,
            round(tr.log_loss, 3) if tr.n_resolved else None,
        ]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(r, j, v if v is not None else "—")
            cell.border = _BORDER
            cell.alignment = Alignment(horizontal="left" if j == 1 else "center", vertical="center")
            if j == 4 and v is not None:
                cell.number_format = "0.0%"
                cell.font = Font(bold=True)
        r += 1

    note = (
        "Nota honesta: el accuracy del Mundial está inflado por la fase de grupos (muchos favoritos "
        "claros) y bajará en eliminatorias. El número fiable es el BRIER (más bajo = mejor). Casi todos "
        "los fallos son EMPATES, que un modelo 1X2 rara vez marca como resultado más probable."
    )
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    nc = ws.cell(r, 1, note)
    nc.font = Font(italic=True, size=9, color="C00000")
    nc.alignment = _WRAP
    ws.row_dimensions[r].height = 46

    # Prueba de transparencia: últimas del Mundial ya resueltas (acierto/fallo).
    r += 2
    ws.cell(r, 1, "Últimas predicciones del Mundial ya resueltas").font = Font(bold=True, size=11)
    r += 1
    sub = ["Resultado", "Local", "Marcador", "Visitante", "Predicción (L/E/V)"]
    for j, h in enumerate(sub, start=1):
        c = ws.cell(r, j, h)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = _CENTER
        c.border = _BORDER
    r += 1
    recent = [v for v in history.recent(competition_code="WC_2026", limit=200) if v.resolved][:16]
    ok_fill = PatternFill("solid", fgColor="C6E0B4")
    miss_fill = PatternFill("solid", fgColor="F8CBAD")
    for v in recent:
        probs = f"{v.prob_home:.0%} / {v.prob_draw:.0%} / {v.prob_away:.0%}"
        cells = [
            "ACIERTO" if v.correct_pick else "FALLO",
            v.home_name,
            f"{v.home_score}-{v.away_score}",
            v.away_name,
            probs,
        ]
        for j, val in enumerate(cells, start=1):
            cell = ws.cell(r, j, val)
            cell.border = _BORDER
            cell.alignment = _CENTER
        ws.cell(r, 1).fill = ok_fill if v.correct_pick else miss_fill
        r += 1


def main() -> None:
    args = sys.argv[1:]
    comp = "WC_2026"
    if "--comp" in args:
        comp = args[args.index("--comp") + 1]
    include_all = "--all" in args  # incluir todas las competiciones, no solo el Mundial

    snap = load_snapshot(default_snapshot_path())
    catalog = InMemoryCatalogRepository(snap.competitions, snap.teams, snap.matches)
    international = {c.code for c in snap.competitions if c.is_international}
    predictor = PredictionService(
        catalog,
        seed_ratings=snap.seed_ratings,
        seed_counts=snap.seed_counts,
        international_competitions=international,
    )

    scheduled = [
        m
        for m in snap.matches
        if m.status == MatchStatus.SCHEDULED and (include_all or m.competition_code == comp)
    ]
    scheduled.sort(key=lambda m: m.kickoff_utc)

    comp_name = next((c.name for c in snap.competitions if c.code == comp), comp)
    print(f"Partidos programados a predecir: {len(scheduled)} ({comp_name})")
    if not scheduled:
        print("No hay partidos programados con equipos definidos. (¿Eliminatorias aún sin definir?)")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Predicciones"
    ncols = len(_COLS)

    # Título + subtítulo + disclaimer
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.cell(1, 1, f"Predicciones — {comp_name}").font = _TITLE_FONT
    ws.cell(1, 1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24

    gen = datetime.now(MX_TZ).strftime("%Y-%m-%d %H:%M")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws.cell(2, 1, f"Generado: {gen} (hora de México) · Modelo: Elo sembrado con historial internacional real").font = _SUB_FONT

    disclaimer = (
        "AVISO: probabilidades, NO certezas. Modelo baseline (Elo) con ratings sembrados de "
        "~8000 partidos internacionales reales. La confianza es provisional. En eliminatorias no hay "
        "empate al 'tiempo regular' como desenlace final, pero el modelo da P(empate) del tiempo "
        "reglamentario. Over/Under y BTTS no se modelan para selecciones (solo ligas de clubes)."
    )
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
    dcell = ws.cell(3, 1, disclaimer)
    dcell.font = Font(italic=True, size=9, color="C00000")
    dcell.alignment = _WRAP
    ws.row_dimensions[3].height = 42

    # Cabecera
    header_row = 5
    for j, (name, width) in enumerate(_COLS, start=1):
        c = ws.cell(header_row, j, name)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = _CENTER
        c.border = _BORDER
        ws.column_dimensions[get_column_letter(j)].width = width

    # Filas
    r = header_row + 1
    n_written = 0
    for m in scheduled:
        try:
            p = predictor.predict_for_match(m)
        except PredictionUnavailable:
            continue

        probs = {"home": p.prob_home, "draw": p.prob_draw, "away": p.prob_away}
        fav = max(probs, key=probs.get)
        if fav == "home":
            pronostico = f"Gana {m.home_team.name}"
        elif fav == "away":
            pronostico = f"Gana {m.away_team.name}"
        else:
            pronostico = "Empate"

        kickoff_mx = m.kickoff_utc.astimezone(MX_TZ).replace(tzinfo=None)
        ctx = " | ".join(n.message for n in p.context) if p.context else ""

        values = [
            kickoff_mx,
            m.stage or "—",
            m.home_team.name,
            m.away_team.name,
            p.prob_home,
            p.prob_draw,
            p.prob_away,
            pronostico,
            p.likely_scoreline,
            round(p.exp_home_goals, 2),
            round(p.exp_away_goals, 2),
            p.confidence_band,
            round(p.rating_home),
            round(p.rating_away),
            ctx,
        ]
        for j, v in enumerate(values, start=1):
            cell = ws.cell(r, j, v)
            cell.border = _BORDER
            cell.alignment = _WRAP if j == ncols else _CENTER
            if j == 1:
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif j in (5, 6, 7):  # probabilidades
                cell.number_format = "0%"
        # Resaltar la probabilidad del resultado más probable.
        fav_col = {"home": 5, "draw": 6, "away": 7}[fav]
        ws.cell(r, fav_col).fill = _FAV_FILL
        ws.cell(r, fav_col).font = Font(bold=True)
        # Color de la banda de confianza.
        band = str(p.confidence_band).lower()
        if band in _BAND_FILL:
            ws.cell(r, 12).fill = _BAND_FILL[band]
        r += 1
        n_written += 1

    ws.freeze_panes = ws.cell(header_row + 1, 1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{r - 1}"

    # Segunda hoja: desempeño / track record (cuántas viene acertando el modelo).
    ws2 = wb.create_sheet("Desempeño")
    _write_track_record_sheet(ws2, catalog, snap, international, predictor)

    out_dir = _desktop_dir()
    fname = f"Predicciones_{comp}_{datetime.now(MX_TZ).strftime('%Y-%m-%d')}.xlsx"
    out_path = out_dir / fname
    wb.save(out_path)
    print(f"OK: {n_written} predicciones escritas.")
    print(f"Excel guardado en: {out_path}")


if __name__ == "__main__":
    main()
